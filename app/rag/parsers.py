from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import time
import xml.etree.ElementTree as ET
import zipfile
import mimetypes

from app.config import get_settings


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    mime_type: str | None
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)
    extraction_ms: float = 0.0
    ocr_ms: float = 0.0
    ocr_pages: int = 0


@dataclass(frozen=True)
class OcrResult:
    text: str
    duration_ms: float
    pages_processed: int = 0


def extract_document_text(path: Path, force_ocr: bool = False) -> ExtractionResult:
    started_at = time.perf_counter()
    mime_type, _ = mimetypes.guess_type(path.name)
    suffix = path.suffix.lower()
    warnings: list[str] = []

    if force_ocr:
        ocr = _extract_with_ocr(path, warnings)
        return _result(
            text=ocr.text,
            mime_type=mime_type,
            ocr_used=True,
            warnings=warnings,
            started_at=started_at,
            ocr=ocr,
        )

    if suffix in {".txt", ".md", ".csv", ".tsv"}:
        return _result(
            text=path.read_text(encoding="utf-8", errors="ignore"),
            mime_type=mime_type,
            warnings=warnings,
            started_at=started_at,
        )

    if suffix == ".pdf":
        text = _extract_pdf_text(path, warnings)
        if _looks_like_scanned_document(text):
            ocr = _extract_with_ocr(path, warnings)
            return _result(
                text=ocr.text,
                mime_type=mime_type,
                ocr_used=True,
                warnings=warnings,
                started_at=started_at,
                ocr=ocr,
            )
        return _result(text=text, mime_type=mime_type, warnings=warnings, started_at=started_at)

    if suffix == ".docx":
        return _result(
            text=_extract_docx_text(path, warnings),
            mime_type=mime_type,
            warnings=warnings,
            started_at=started_at,
        )

    if suffix == ".xlsx":
        return _result(
            text=_extract_xlsx_text(path, warnings),
            mime_type=mime_type,
            warnings=warnings,
            started_at=started_at,
        )

    if suffix == ".pptx":
        return _result(
            text=_extract_pptx_text(path, warnings),
            mime_type=mime_type,
            warnings=warnings,
            started_at=started_at,
        )

    warnings.append(f"No parser configured for {suffix}; attempted utf-8 text fallback.")
    return _result(
        text=path.read_text(encoding="utf-8", errors="ignore"),
        mime_type=mime_type,
        warnings=warnings,
        started_at=started_at,
    )


def _result(
    *,
    text: str,
    mime_type: str | None,
    warnings: list[str],
    started_at: float,
    ocr_used: bool = False,
    ocr: OcrResult | None = None,
) -> ExtractionResult:
    return ExtractionResult(
        text=text,
        mime_type=mime_type,
        ocr_used=ocr_used,
        warnings=warnings,
        extraction_ms=_elapsed_ms(started_at),
        ocr_ms=ocr.duration_ms if ocr else 0.0,
        ocr_pages=ocr.pages_processed if ocr else 0,
    )


def _extract_pdf_text(path: Path, warnings: list[str]) -> str:
    try:
        import fitz  # type: ignore
    except ImportError:
        warnings.append("PyMuPDF is not installed; attempted utf-8 text fallback for PDF.")
        return path.read_text(encoding="utf-8", errors="ignore")

    pages: list[str] = []
    with fitz.open(path) as document:
        for page in document:
            pages.append(page.get_text("text"))
    return "\n\n".join(pages)


def _extract_docx_text(path: Path, warnings: list[str]) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        warnings.append("python-docx is not installed; attempted utf-8 text fallback for DOCX.")
        return path.read_text(encoding="utf-8", errors="ignore")

    document = docx.Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def _extract_xlsx_text(path: Path, warnings: list[str]) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        warnings.append("openpyxl is not installed; attempted utf-8 text fallback for XLSX.")
        return path.read_text(encoding="utf-8", errors="ignore")

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[str] = []
    for worksheet in workbook.worksheets:
        rows.append(f"Sheet: {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
    workbook.close()
    return "\n".join(rows)


def _extract_pptx_text(path: Path, warnings: list[str]) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            slide_names = sorted(
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
            slides: list[str] = []
            namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for index, slide_name in enumerate(slide_names, start=1):
                root = ET.fromstring(archive.read(slide_name))
                texts = [
                    node.text.strip()
                    for node in root.findall(".//a:t", namespace)
                    if node.text and node.text.strip()
                ]
                if texts:
                    slides.append(f"Slide {index}: " + " ".join(texts))
            return "\n".join(slides)
    except (zipfile.BadZipFile, ET.ParseError):
        warnings.append("Unable to parse PPTX; attempted utf-8 text fallback.")
        return path.read_text(encoding="utf-8", errors="ignore")


def _extract_with_ocr(path: Path, warnings: list[str]) -> OcrResult:
    started_at = time.perf_counter()
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except ImportError:
        warnings.append("OCR dependencies are not installed; returning empty OCR result.")
        return OcrResult(text="", duration_ms=_elapsed_ms(started_at))

    settings = get_settings()
    language = settings.ocr_language.strip() or None

    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image, lang=language)
        return OcrResult(text=text, duration_ms=_elapsed_ms(started_at), pages_processed=1)

    if path.suffix.lower() == ".pdf":
        text, pages_processed = _extract_pdf_with_ocr(path, warnings, pytesseract, Image)
        return OcrResult(
            text=text,
            duration_ms=_elapsed_ms(started_at),
            pages_processed=pages_processed,
        )

    warnings.append(f"OCR is not configured for {path.suffix.lower() or 'files without extension'}.")
    return OcrResult(text="", duration_ms=_elapsed_ms(started_at))


def _extract_pdf_with_ocr(path: Path, warnings: list[str], pytesseract, Image) -> tuple[str, int]:
    try:
        import fitz  # type: ignore
    except ImportError:
        warnings.append("PyMuPDF is not installed; returning empty PDF OCR result.")
        return "", 0

    settings = get_settings()
    language = settings.ocr_language.strip() or None
    max_pages = max(settings.ocr_max_pdf_pages, 1)
    dpi = max(settings.ocr_pdf_dpi, 72)
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    pages: list[str] = []
    pages_processed = 0

    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            if page_index >= max_pages:
                warnings.append(
                    f"PDF OCR stopped after {max_pages} pages; increase OCR_MAX_PDF_PAGES to scan more."
                )
                break
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            pages_processed += 1
            with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                page_text = pytesseract.image_to_string(image, lang=language).strip()
            if page_text:
                pages.append(page_text)

    if not pages:
        warnings.append("PDF OCR produced no text.")
    return "\n\n".join(pages), pages_processed


def _looks_like_scanned_document(text: str) -> bool:
    return len(text.strip()) < 100


def _elapsed_ms(started_at: float) -> float:
    return round((time.perf_counter() - started_at) * 1000, 3)
